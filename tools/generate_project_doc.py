from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Project_Documentation_PersonalLedger.docx"

ACCENT = RGBColor(30, 76, 138)
ACCENT_LIGHT = "DCE6F2"
TEXT_DARK = RGBColor(34, 34, 34)
TEXT_MUTED = RGBColor(90, 90, 90)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_page_margins(section):
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def set_default_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT_DARK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for style_name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5)]:
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.color.rgb = ACCENT
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = TEXT_MUTED
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_run(paragraph, text, *, bold=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return run


def add_cover(document):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(110)
    add_run(p, "PersonalLedger", bold=True, size=24, color=ACCENT)

    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p2, "项目文档", bold=True, size=20, color=TEXT_DARK)

    p3 = document.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.space_before = Pt(8)
    add_run(
        p3,
        "面向 Android 本地记账场景的个人账本应用说明书",
        size=11,
        color=TEXT_MUTED,
    )

    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Cm(4)
    table.columns[1].width = Cm(10)
    rows = [
        ("应用名称", "轻账 / PersonalLedger"),
        ("当前版本", "1.2.2"),
        ("技术栈", "Kotlin、AndroidX、DataStore、MPAndroidChart、Material Components"),
        ("文档用途", "课程项目、毕业设计归档、作品集展示"),
    ]
    for i, (k, v) in enumerate(rows):
        left = table.cell(i, 0)
        right = table.cell(i, 1)
        left.text = k
        right.text = v
        left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(left, ACCENT_LIGHT)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_info_table(document):
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["维度", "内容", "维度", "内容"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        set_cell_shading(cell, ACCENT_LIGHT)
    values = [
        ("开发语言", "Kotlin", "最低版本", "Android 7.0 / API 24"),
        ("目标版本", "Android API 36", "Java 版本", "11"),
        ("构建工具", "Gradle Kotlin DSL", "包名", "com.example.personalledger"),
    ]
    for row in values:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_feature_table(document):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["模块", "核心能力", "实现要点"]):
        cell.text = text
        set_cell_shading(cell, ACCENT_LIGHT)
    rows = [
        ("首页记账", "展示本月结余、今日收支、预算进度与最近记录", "HomeFragment + RecyclerView + LiveData"),
        ("账单录入", "通过底部弹窗录入收入/支出、分类、金额与备注", "BottomSheetDialogFragment + 表单校验"),
        ("预算管理", "设置月预算并根据当月支出计算进度与预警", "DataStore 持久化 + 进度条颜色反馈"),
        ("报表分析", "饼图、柱状图、折线图展示消费结构与资产趋势", "MPAndroidChart + ViewModel 数据聚合"),
        ("数据导出", "导出 CSV 并调起系统分享能力", "Storage Access Framework + Intent 分享"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_data_table(document):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["字段", "类型", "说明"]):
        cell.text = text
        set_cell_shading(cell, ACCENT_LIGHT)
    rows = [
        ("amount", "String", "账单金额，当前以字符串形式保存并在业务层转为数值"),
        ("note", "String", "备注信息"),
        ("time", "String", "记录时间，格式为 yyyy-MM-dd HH:mm"),
        ("isExpense", "Boolean", "是否为支出"),
        ("categoryName", "String", "分类名称"),
        ("categoryIconRes", "Int", "分类图标资源 ID"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        add_run(p, item)


def add_numbered(document, items):
    for item in items:
        p = document.add_paragraph(style="List Number")
        add_run(p, item)


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "PersonalLedger 项目文档", size=9, color=TEXT_MUTED)


def build():
    document = Document()
    set_page_margins(document.sections[0])
    set_default_styles(document)
    add_footer(document.sections[0])
    add_cover(document)

    section = document.sections[-1]
    set_page_margins(section)

    document.add_heading("1. 项目概述", level=1)
    p = document.add_paragraph()
    add_run(
        p,
        "PersonalLedger（应用名“轻账”）是一款基于 Android 原生开发的个人记账应用，"
        "面向日常收支记录、预算控制与数据复盘场景。项目以本地离线使用为核心，"
        "强调录入效率、统计直观性与轻量级实现。"
    )
    add_info_table(document)

    document.add_heading("2. 建设目标", level=1)
    add_bullets(
        document,
        [
            "提供简洁直观的个人记账体验，降低日常记账门槛。",
            "支持收入与支出双类型录入，并配套分类管理能力。",
            "通过预算设置和图表展示提升用户对消费结构的感知能力。",
            "支持 CSV 导出与分享，便于备份和二次分析。",
        ],
    )

    document.add_heading("3. 功能模块说明", level=1)
    add_feature_table(document)

    document.add_heading("4. 技术方案与架构", level=1)
    add_bullets(
        document,
        [
            "UI 层由 MainActivity、HomeFragment、ReportFragment、AddEntryBottomSheetDialogFragment 组成，负责页面承载与交互。",
            "状态与业务层由 MainViewModel 统一协调，使用 LiveData 向界面发布数据变化。",
            "数据持久化层使用 DataStoreManager，将累计资产、预算金额和历史账单列表保存到 Jetpack DataStore Preferences。",
            "图表展示使用 MPAndroidChart，分别输出支出分类占比、近 7 日支出趋势和累计结余变化。",
        ],
    )
    p = document.add_paragraph()
    add_run(p, "架构流转：", bold=True)
    add_run(p, " 界面事件 -> ViewModel 业务处理 -> DataStore 持久化 -> LiveData 回流刷新 UI。")

    document.add_heading("5. 页面与交互设计", level=1)
    add_bullets(
        document,
        [
            "主界面采用底部导航，在“记账”和“报表”两个核心页面间切换。",
            "首页聚合展示本月结余、今日收支、预算使用率与最近账单，兼顾总览与快速操作。",
            "账单录入使用 BottomSheet 形式，符合移动端单手操作习惯。",
            "报表页结合卡片容器与图表组件，增强数据可读性和信息层次。",
            "支持长按删除、导出分享等轻量交互，避免页面过度复杂化。",
        ],
    )

    document.add_heading("6. 数据设计", level=1)
    p = document.add_paragraph()
    add_run(p, "核心实体 LedgerItem", bold=True)
    add_data_table(document)
    p = document.add_paragraph()
    add_run(
        p,
        "本项目还维护了预算金额和累计资产等关键数据。其中历史账单列表以 JSON 字符串形式保存，"
        "并保留旧字段兼容逻辑，以适配历史版本数据迁移。"
    )

    document.add_heading("7. 核心业务逻辑", level=1)
    add_bullets(
        document,
        [
            "新增支出时累计资产减少，新增收入时累计资产增加；删除或编辑账单时会同步回滚并重算。",
            "今日收支和当月支出基于账单时间前缀进行筛选聚合，格式分别使用 yyyy-MM-dd 与 yyyy-MM。",
            "预算进度按照“当月支出 / 当前预算 * 100%”计算，并将结果限制在 0% 到 100% 区间。",
            "报表数据在 ViewModel 中按分类、按日期和按时间顺序进行整理，再转换为图表数据集。",
        ],
    )

    document.add_heading("8. 项目结构", level=1)
    add_bullets(
        document,
        [
            "app/src/main/java/com/example/personalledger：核心 Kotlin 业务与界面代码。",
            "app/src/main/res/layout：主页面、报表页、底部录入面板与列表项布局。",
            "app/src/main/res/drawable、font、anim、menu：图标、主题背景、字体、动画与菜单资源。",
            "build.gradle.kts、settings.gradle.kts：Gradle 构建配置。",
        ],
    )

    document.add_heading("9. 构建与运行", level=1)
    add_numbered(
        document,
        [
            "使用 Android Studio 打开项目根目录。",
            "等待 Gradle 同步完成并准备 Android SDK。",
            "连接真机或启动模拟器后运行 app 模块。",
            "如需构建安装包，可执行 .\\gradlew assembleRelease。",
        ],
    )

    document.add_heading("10. 项目亮点", level=1)
    add_bullets(
        document,
        [
            "功能闭环完整，覆盖记账、预算、统计、导出与分享。",
            "使用 DataStore 替代 SharedPreferences，体现较新的本地存储实践。",
            "通过 MPAndroidChart 增强数据可视化效果，便于展示作品完整度。",
            "采用 ViewBinding 和较清晰的页面分层，降低常见空指针与耦合风险。",
        ],
    )

    document.add_heading("11. 当前限制与优化方向", level=1)
    add_bullets(
        document,
        [
            "当前数据完全保存在本地，不支持账号体系与云同步。",
            "账单列表以 JSON 整体存储，面对大规模数据时可维护性和查询能力有限。",
            "暂未提供搜索、筛选、按月查看和账单编辑历史等高级能力。",
            "自动化测试覆盖较少，后续可补充单元测试与 UI 测试。",
            "后续可考虑引入 Room、通知提醒、备份恢复和多账本支持。",
        ],
    )

    document.add_heading("12. 适用场景与结论", level=1)
    p = document.add_paragraph()
    add_run(
        p,
        "PersonalLedger 适合作为 Android 原生课程设计、毕业设计作品或个人作品集项目。"
        "它在有限复杂度下实现了较完整的记账产品闭环，既具备实际演示价值，也为后续扩展为"
        "更成熟的个人财务管理应用提供了清晰基础。"
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
